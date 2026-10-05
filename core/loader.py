"""
Dataset loading for files on disk and uploaded bytes.

Supports CSV, TSV, Excel, JSON and Parquet. After reading, columns are
lightly normalised so downstream inference sees proper dtypes:

- column names are stripped of surrounding whitespace
- text columns holding numbers become numeric
- text columns holding dates become datetimes
"""

from __future__ import annotations

import io
import re
import warnings
from pathlib import Path

import pandas as pd


SUPPORTED_EXTENSIONS = {
    ".csv", ".tsv", ".txt", ".xlsx", ".xls", ".json", ".parquet",
}

EXCEL_EXTENSIONS = {".xlsx", ".xls"}

# Values like "12:30" are times of day, not dates.
_TIME_OF_DAY = re.compile(r"^\s*\d{1,2}:\d{2}(:\d{2})?\s*$")
_PURE_NUMBER = re.compile(r"^\s*-?\d+(\.\d+)?\s*$")
# A date has a separator or a month name; codes like "ST001" have neither.
_DATE_SHAPE = re.compile(
    r"[-/.,: ]|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec",
    re.IGNORECASE,
)
# Parsed years outside this range mean the values were not dates.
_PLAUSIBLE_YEARS = (1800, 2200)


def load_dataset(file_path, sheet_name=None) -> pd.DataFrame:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    with open(path, "rb") as handle:
        content = handle.read()

    return load_dataset_from_bytes(content, path.name, sheet_name)


def load_dataset_from_bytes(
    content: bytes,
    filename: str,
    sheet_name: str | None = None,
) -> pd.DataFrame:

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: '{extension}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    buffer = io.BytesIO(content)

    if extension in EXCEL_EXTENSIONS:
        if sheet_name is None:
            sheet_name = default_sheet(content, filename)
        dataframe = pd.read_excel(buffer, sheet_name=sheet_name)

    elif extension in {".csv", ".txt"}:
        dataframe = pd.read_csv(buffer)

        # Fall back to delimiter sniffing for ';' or '|' separated files.
        if dataframe.shape[1] == 1:
            buffer.seek(0)
            dataframe = pd.read_csv(buffer, sep=None, engine="python")

    elif extension == ".tsv":
        dataframe = pd.read_csv(buffer, sep="\t")

    elif extension == ".json":
        dataframe = pd.read_json(buffer)

    else:
        dataframe = pd.read_parquet(buffer)

    if dataframe.empty:
        raise ValueError("Dataset is empty.")

    return normalize_dataframe(dataframe)


def load_all_sheets(content: bytes, filename: str) -> dict[str, pd.DataFrame]:
    """
    Every non-empty sheet of an Excel file, read in one pass and
    normalised. Other formats return their single table under "".
    """

    if Path(filename).suffix.lower() not in EXCEL_EXTENSIONS:
        return {"": load_dataset_from_bytes(content, filename)}

    frames = pd.read_excel(io.BytesIO(content), sheet_name=None)

    return {
        str(name): normalize_dataframe(frame)
        for name, frame in frames.items()
        if not frame.empty
    }


def list_sheets(content: bytes, filename: str) -> list[str]:
    """Sheet names for Excel files; empty list for other formats."""

    if Path(filename).suffix.lower() not in EXCEL_EXTENSIONS:
        return []

    with pd.ExcelFile(io.BytesIO(content)) as workbook:
        return [str(name) for name in workbook.sheet_names]


def default_sheet(content: bytes, filename: str) -> str | None:
    """
    The widest sheet (most columns, then most rows) is assumed to hold
    the main analysis table. Lookup / reference sheets are usually
    narrow, while modelling or training tables carry the most features.
    """

    sheets = list_sheets(content, filename)

    if not sheets:
        return None

    sizes = _sheet_sizes_from_metadata(content, filename)

    if sizes is None:
        sizes = {}
        with pd.ExcelFile(io.BytesIO(content)) as workbook:
            for sheet in sheets:
                frame = workbook.parse(sheet)
                sizes[sheet] = (frame.shape[1], frame.shape[0])

    return max(sheets, key=lambda sheet: sizes.get(sheet, (0, 0)))


def _sheet_sizes_from_metadata(
    content: bytes,
    filename: str,
) -> dict[str, tuple[int, int]] | None:
    """Read sheet dimensions without parsing cells (.xlsx only)."""

    if Path(filename).suffix.lower() != ".xlsx":
        return None

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True,
        )
    except Exception:
        return None

    try:
        sizes = {}

        for sheet in workbook.worksheets:
            if sheet.max_row is None or sheet.max_column is None:
                return None
            sizes[sheet.title] = (sheet.max_column, sheet.max_row)

        return sizes
    finally:
        workbook.close()


# ----------------------------------------------------------------------
# Normalisation
# ----------------------------------------------------------------------

def normalize_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    result = dataframe.copy()

    result.columns = _unique_names(
        [str(column).strip() for column in result.columns]
    )

    result = result.dropna(axis=0, how="all")

    for column in result.columns:
        series = result[column]

        if not _is_text(series):
            continue

        converted = _try_numeric(series)

        if converted is None:
            converted = _try_datetime(series)

        if converted is not None:
            result[column] = converted

    return result.reset_index(drop=True)


def _is_text(series: pd.Series) -> bool:
    return (
        pd.api.types.is_object_dtype(series)
        or pd.api.types.is_string_dtype(series)
    )


def _try_numeric(series: pd.Series) -> pd.Series | None:
    non_null = series.dropna().astype(str).str.strip()

    if non_null.empty:
        return None

    cleaned = non_null.str.replace(",", "", regex=False)

    if cleaned.str.match(_PURE_NUMBER).mean() < 0.98:
        return None

    return pd.to_numeric(
        series.astype(str).str.strip().str.replace(",", "", regex=False),
        errors="coerce",
    )


def _try_datetime(series: pd.Series) -> pd.Series | None:
    non_null = series.dropna().astype(str).str.strip()

    if non_null.empty:
        return None

    sample = non_null.sample(
        min(len(non_null), 300),
        random_state=0,
    )

    if sample.str.match(_TIME_OF_DAY).mean() > 0.5:
        return None

    if sample.str.match(_PURE_NUMBER).mean() > 0.5:
        return None

    # Dates contain at least one digit and one separator or month name.
    if sample.str.contains(r"\d").mean() < 0.9:
        return None

    if sample.str.contains(_DATE_SHAPE).mean() < 0.9:
        return None

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        parsed_sample = pd.to_datetime(
            sample,
            errors="coerce",
            format="mixed",
        )

    if parsed_sample.notna().mean() < 0.95:
        return None

    years = parsed_sample.dropna().dt.year
    low, high = _PLAUSIBLE_YEARS
    if years.between(low, high).mean() < 0.95:
        return None

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return pd.to_datetime(
            series,
            errors="coerce",
            format="mixed",
        )


def _unique_names(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []

    for name in names:
        name = name or "Unnamed"

        if name in seen:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            result.append(name)

    return result
